import asyncio
import copy
import datetime
import importlib
import unittest
from types import SimpleNamespace
from unittest.mock import AsyncMock, patch

from bson import ObjectId
from openpyxl import load_workbook


orders_module = importlib.import_module("zns-chatbot.plugins.orders")


class _Result:
    def __init__(self, matched_count=0, modified_count=None, deleted_count=None):
        self.matched_count = matched_count
        self.modified_count = matched_count if modified_count is None else modified_count
        self.deleted_count = matched_count if deleted_count is None else deleted_count


class _Cursor:
    def __init__(self, documents):
        self.documents = documents

    async def to_list(self, _length):
        return copy.deepcopy(self.documents)

    def sort(self, key, direction=1):
        self.documents.sort(
            key=lambda document: document.get(key), reverse=direction < 0
        )
        return self

    def __aiter__(self):
        self._iterator = iter(copy.deepcopy(self.documents))
        return self

    async def __anext__(self):
        try:
            return next(self._iterator)
        except StopIteration as error:
            raise StopAsyncIteration from error


def _matches(document, query):
    for key, expected in query.items():
        if key == "$or":
            if not any(_matches(document, clause) for clause in expected):
                return False
            continue
        if key == "$and":
            if not all(_matches(document, clause) for clause in expected):
                return False
            continue
        value = document
        exists = True
        for part in key.split("."):
            if not isinstance(value, dict) or part not in value:
                exists = False
                break
            value = value[part]
        if isinstance(expected, dict) and any(
            operator.startswith("$") for operator in expected
        ):
            if "$exists" in expected and exists != expected["$exists"]:
                return False
            if "$ne" in expected and exists and value == expected["$ne"]:
                return False
            if "$in" in expected and (not exists or value not in expected["$in"]):
                return False
            if "$nin" in expected and exists and value in expected["$nin"]:
                return False
            if "$lte" in expected and (not exists or value > expected["$lte"]):
                return False
            if "$gt" in expected and (not exists or value <= expected["$gt"]):
                return False
        elif not exists or value != expected:
            return False
    return True


def _project(document, projection):
    if projection is None:
        return copy.deepcopy(document)
    result = {}
    if projection.get("_id", 1) and "_id" in document:
        result["_id"] = document["_id"]
    for key, include in projection.items():
        if include and key in document:
            result[key] = document[key]
    return copy.deepcopy(result)


class _Collection:
    def __init__(self, documents=None):
        self.documents = [copy.deepcopy(document) for document in documents or []]
        self.lock = asyncio.Lock()

    async def create_index(self, *_args, **_kwargs):
        return "test-index"

    async def insert_one(self, document):
        async with self.lock:
            self.documents.append(copy.deepcopy(document))
        return _Result(1)

    def find(self, query, projection=None):
        return _Cursor([
            _project(document, projection)
            for document in self.documents
            if _matches(document, query)
        ])

    async def find_one(self, query, projection=None):
        async with self.lock:
            for document in self.documents:
                if _matches(document, query):
                    return _project(document, projection)
        return None

    async def update_one(self, query, update, upsert=False):
        async with self.lock:
            for document in self.documents:
                if _matches(document, query):
                    self._apply_update(document, update)
                    return _Result(1)
            if upsert:
                document = {
                    key: value for key, value in query.items()
                    if not isinstance(value, dict)
                }
                self._apply_update(document, update)
                self.documents.append(document)
                return _Result(0)
        return _Result(0)

    async def delete_one(self, query):
        async with self.lock:
            for index, document in enumerate(self.documents):
                if _matches(document, query):
                    self.documents.pop(index)
                    return _Result(1, deleted_count=1)
        return _Result(0, deleted_count=0)

    async def find_one_and_update(self, query, update, sort=None):
        async with self.lock:
            matches = [document for document in self.documents if _matches(document, query)]
            if not matches:
                return None
            if sort:
                key, direction = sort[0]
                matches.sort(key=lambda document: document[key], reverse=direction < 0)
            document = matches[0]
            original = copy.deepcopy(document)
            self._apply_update(document, update)
            return original

    @staticmethod
    def _apply_update(document, update):
        document.update(update.get("$setOnInsert", {}))
        for key, value in update.get("$set", {}).items():
            target = document
            parts = key.split(".")
            for part in parts[:-1]:
                target = target.setdefault(part, {})
            target[parts[-1]] = value
        for key in update.get("$unset", {}):
            target = document
            parts = key.split(".")
            for part in parts[:-1]:
                target = target.get(part, {})
            target.pop(parts[-1], None)


def _orders_service(event_key="grodno_26"):
    service = orders_module.Orders.__new__(orders_module.Orders)
    service.config = SimpleNamespace(
        orders=SimpleNamespace(
            event_key=event_key,
            payment_reminder_after=datetime.timedelta(days=2),
            notification_check_interval=datetime.timedelta(minutes=5),
        ),
    )
    service.food_db = _Collection()
    service.capacity_db = _Collection()
    service._capacity_slots_ready = set()
    service._capacity_slots_lock = asyncio.Lock()
    service.menu = service.get_menu()
    return service


class CapacityTestCase(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        super().setUp()
        # Keep capacity scenarios independent of production event limits.
        self.capacity_limits = {
            orders_module.SHUTTLE_SERVICE: 3,
            orders_module.GRODNO_OVERVIEW_SERVICE: 4,
            orders_module.GRODNO_GORODNITSA_SERVICE: 5,
        }
        self.shuttle_capacity = self.capacity_limits[orders_module.SHUTTLE_SERVICE]
        self.enterContext(patch.dict(
            orders_module.CAPACITY_LIMITS, self.capacity_limits, clear=True
        ))


class ShuttleCapacityTests(CapacityTestCase):
    async def test_deleted_order_frees_capacity_without_restart(self):
        for check_availability in (True, False):
            with self.subTest(check_availability=check_availability):
                service = _orders_service()
                old_order = {
                    "_id": ObjectId(),
                    "event_key": "grodno_26",
                    "proof_file": "proof.jpg",
                    "proof_received": datetime.datetime(2026, 1, 1),
                    "choice": {"extras": {"shuttle": 65}},
                }
                await service.food_db.insert_one(old_order)
                with patch.dict(orders_module.CAPACITY_LIMITS, {"shuttle": 1}):
                    await service._ensure_capacity_slots("shuttle")
                    self.assertFalse(await service.shuttle_available())
                    await service.food_db.delete_one({"_id": old_order["_id"]})
                    if check_availability:
                        self.assertTrue(await service.shuttle_available())
                    new_order = {
                        **old_order,
                        "_id": ObjectId(),
                        "proof_received": datetime.datetime(2026, 1, 2),
                    }
                    await service.food_db.insert_one(new_order)
                    self.assertTrue(await service.reserve_service_seat(
                        "shuttle", new_order["_id"], new_order
                    ))
                    self.assertFalse(await service.shuttle_available())
                    self.assertIsNone(await service.capacity_db.find_one({
                        "reservation_id": old_order["_id"],
                    }))

    def test_capacity_uses_nested_orders_config(self):
        service = _orders_service(event_key="grodno_30")

        self.assertEqual(service._capacity_event_key(), "grodno_30")

    async def test_concurrent_reservations_respect_shuttle_capacity(self):
        service = _orders_service()
        order_ids = [ObjectId() for _ in range(self.shuttle_capacity + 1)]

        results = await asyncio.gather(*(
            service.reserve_shuttle_seat(order_id) for order_id in order_ids
        ))

        self.assertEqual(sum(results), self.shuttle_capacity)
        self.assertFalse(await service.shuttle_available())

    async def test_releasing_a_seat_allows_another_order(self):
        service = _orders_service()
        order_ids = [ObjectId() for _ in range(self.shuttle_capacity)]
        for order_id in order_ids:
            self.assertTrue(await service.reserve_shuttle_seat(order_id))

        self.assertFalse(await service.reserve_shuttle_seat(ObjectId()))
        await service.release_shuttle_seat(order_ids[0])
        self.assertTrue(await service.reserve_shuttle_seat(ObjectId()))

    async def test_only_paid_existing_transfer_can_still_open_when_full(self):
        service = _orders_service()
        for _ in range(self.shuttle_capacity):
            self.assertTrue(await service.reserve_shuttle_seat(ObjectId()))

        choice = {"extras": {"shuttle": {"count": 1}}}
        self.assertFalse(await service.shuttle_available(choice))
        self.assertTrue(await service.shuttle_available(choice, True))
        self.assertTrue(orders_module.choice_has_shuttle(choice))

    async def test_unpaid_order_does_not_reserve_a_transfer_seat(self):
        service = _orders_service()
        update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        update.base = service
        update.user = 123
        update.config = SimpleNamespace(event_key="grodno_26")
        update.handle_cq_start = AsyncMock()

        await update.create_order({
            "total": 65,
            "extras": {"total": 65, "shuttle": 65},
        })

        reserved = [
            slot for slot in service.capacity_db.documents
            if "reservation_id" in slot
        ]
        self.assertEqual(reserved, [])
        self.assertTrue(await service.shuttle_available())

    async def test_full_transfer_rejects_order_without_saving_it(self):
        service = _orders_service()
        for _ in range(self.shuttle_capacity):
            self.assertTrue(await service.reserve_shuttle_seat(ObjectId()))
        update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        update.base = service
        update.user = 123
        update.config = SimpleNamespace(event_key="grodno_26")
        choice = {"extras": {"shuttle": 65}}

        with self.assertRaises(orders_module.ShuttleFullError):
            await update.create_order(choice)

        self.assertEqual(service.food_db.documents, [])


class PaidOrderCapacityTests(CapacityTestCase):
    @staticmethod
    def _notification_update(reply):
        return SimpleNamespace(
            l=lambda key, **kwargs: f"{key}:{kwargs}",
            update=SimpleNamespace(reply=reply),
        )

    async def test_submitted_proof_reserves_a_seat_before_admin_confirmation(self):
        service = _orders_service()
        order_id = ObjectId()
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "created_at": datetime.datetime(2026, 8, 20),
            "proof_file": "telegram-file",
            "proof_received": datetime.datetime(2026, 8, 21),
            "choice": {
                "total": 65,
                "extras": {"total": 65, "shuttle": 65},
            },
        })

        order, removed = await service.activate_paid_order_capacity(order_id)

        self.assertEqual(removed, [])
        self.assertIn("shuttle", order["choice"]["extras"])
        reservation = await service.capacity_db.find_one({
            "event_key": "grodno_26",
            "service": "shuttle",
            "reservation_id": order_id,
        })
        self.assertIsNotNone(reservation)

    async def test_legacy_cash_marker_does_not_reserve_capacity(self):
        service = _orders_service()
        order_id = ObjectId()
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "proof_file": "cash",
            "choice": {
                "total": 65,
                "extras": {"total": 65, "shuttle": 65},
            },
        })

        order, removed = await service.activate_paid_order_capacity(order_id)

        self.assertFalse(orders_module.order_has_payment_proof(order))
        self.assertEqual(removed, [])
        self.assertEqual(
            [slot for slot in service.capacity_db.documents if "reservation_id" in slot],
            [],
        )

    def test_recalculation_uses_authoritative_service_price(self):
        choice, removed = orders_module.choice_without_capacity_services(
            {
                "total": 1000,
                "extras": {"total": 65, "shuttle": -1000},
            },
            {"shuttle"},
            _orders_service().menu,
        )

        self.assertEqual(removed, ["shuttle"])
        self.assertEqual(choice["total"], 0)
        self.assertEqual(choice["extras"]["total"], 0)

    async def test_capacity_rebuild_preserves_payment_proof_order(self):
        service = _orders_service()
        paid_orders = []
        start = datetime.datetime(2026, 8, 20, 12, 0)
        for index in range(self.shuttle_capacity + 1):
            order = {
                "_id": ObjectId(),
                "user_id": index,
                "event_key": "grodno_26",
                "created_at": start - datetime.timedelta(days=1),
                "proof_file": f"proof-{index}",
                "proof_received": start + datetime.timedelta(minutes=index),
                "choice": {
                    "total": 65,
                    "extras": {"total": 65, "shuttle": 65},
                },
            }
            paid_orders.append(order)
            service.food_db.documents.append(order)

        await service._ensure_capacity_slots("shuttle")

        reserved_ids = {
            slot["reservation_id"]
            for slot in service.capacity_db.documents
            if slot.get("service") == "shuttle" and "reservation_id" in slot
        }
        self.assertEqual(
            reserved_ids,
            {order["_id"] for order in paid_orders[:-1]},
        )

    async def test_capacity_rebuild_displaces_late_reservation_for_earlier_proof(self):
        service = _orders_service()
        proof_time = datetime.datetime(2026, 8, 20, 12, 0)
        orders = []
        for index in range(self.shuttle_capacity + 1):
            order = {
                "_id": ObjectId(),
                "user_id": index,
                "event_key": "grodno_26",
                "created_at": proof_time - datetime.timedelta(days=1),
                "proof_file": f"proof-{index}",
                "proof_received": proof_time + datetime.timedelta(minutes=index),
                "choice": {
                    "total": 65,
                    "extras": {"total": 65, "shuttle": 65},
                },
            }
            orders.append(order)
            service.food_db.documents.append(order)
        for seat, order in enumerate(orders[1:]):
            service.capacity_db.documents.append({
                "_id": f"grodno_26:shuttle:{seat}",
                "event_key": "grodno_26",
                "service": "shuttle",
                "seat": seat,
                "reservation_id": order["_id"],
            })

        await service._ensure_capacity_slots("shuttle")

        reserved_ids = {
            slot["reservation_id"]
            for slot in service.capacity_db.documents
            if "reservation_id" in slot
        }
        self.assertEqual(reserved_ids, {order["_id"] for order in orders[:-1]})
        self.assertNotIn(orders[-1]["_id"], reserved_ids)

    async def test_capacity_rebuild_uses_cash_request_time_not_validation_time(self):
        service = _orders_service()
        cash_id = ObjectId()
        proof_id = ObjectId()
        cash_requested = datetime.datetime(2026, 8, 23, 10, 0)
        service.food_db.documents.extend([
            {
                "_id": cash_id,
                "user_id": 1,
                "event_key": "grodno_26",
                "created_at": cash_requested - datetime.timedelta(days=1),
                "validation": True,
                "validated_at": cash_requested + datetime.timedelta(hours=2),
                "payment_attempt_token": "cash",
                "payment_attempt_created_at": cash_requested,
                "choice": {
                    "total": 65,
                    "extras": {"total": 65, "shuttle": 65},
                },
            },
            {
                "_id": proof_id,
                "user_id": 2,
                "event_key": "grodno_26",
                "created_at": cash_requested - datetime.timedelta(days=1),
                "proof_file": "proof",
                "proof_received": cash_requested + datetime.timedelta(hours=1),
                "choice": {
                    "total": 65,
                    "extras": {"total": 65, "shuttle": 65},
                },
            },
        ])

        with patch.dict(orders_module.CAPACITY_LIMITS, {"shuttle": 1}):
            await service._ensure_capacity_slots("shuttle")

        slot = await service.capacity_db.find_one({
            "event_key": "grodno_26", "service": "shuttle", "seat": 0
        })
        self.assertEqual(slot["reservation_id"], cash_id)

    async def test_stale_unpaid_reconciliation_cannot_strip_paid_order(self):
        service = _orders_service()
        order_id = ObjectId()
        stale_order = {
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "choice": {
                "total": 165,
                "extras": {"total": 65, "shuttle": 65},
            },
        }
        paid_order = copy.deepcopy(stale_order)
        paid_order["proof_file"] = "proof"
        service.food_db.documents.append(paid_order)

        _order, removed = await service.remove_capacity_services_from_order(
            stale_order, {"shuttle"}, reason="unpaid"
        )

        current = await service.food_db.find_one({"_id": order_id})
        self.assertEqual(removed, [])
        self.assertIn("shuttle", current["choice"]["extras"])

    async def test_proof_submission_wins_race_with_stale_edit(self):
        class ProofRaceCollection(_Collection):
            async def update_one(self, query, update, upsert=False):
                if "choice" in update.get("$set", {}):
                    self.documents[0]["proof_file"] = "proof"
                return await super().update_one(query, update, upsert)

        service = _orders_service()
        order_id = ObjectId()
        old_choice = {"total": 10, "extras": {}}
        service.food_db = ProofRaceCollection([{
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "choice": old_choice,
        }])
        update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        update.base = service
        update.user = 123
        update.config = SimpleNamespace(event_key="grodno_26")

        with self.assertRaises(ValueError):
            await update.set_choice(order_id, {"total": 20, "extras": {}})

        current = await service.food_db.find_one({"_id": order_id})
        self.assertEqual(current["choice"], old_choice)
        self.assertTrue(orders_module.order_has_payment_proof(current))

    async def test_validated_order_cannot_cancel_proof_or_release_seat(self):
        service = _orders_service()
        order_id = ObjectId()
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "proof_file": "proof",
            "validation": True,
            "choice": {
                "total": 65,
                "extras": {"total": 65, "shuttle": 65},
            },
        })
        await service.activate_paid_order_capacity(order_id)
        update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        update.base = service
        update.user = 123
        update.config = SimpleNamespace(event_key="grodno_26")
        update.l = lambda key, **_kwargs: key
        update.update = SimpleNamespace(edit_or_reply=AsyncMock())

        await update.handle_cq_pcancel(str(order_id))

        current = await service.food_db.find_one({"_id": order_id})
        self.assertEqual(current["proof_file"], "proof")
        reservation = await service.capacity_db.find_one({
            "service": "shuttle", "reservation_id": order_id
        })
        self.assertIsNotNone(reservation)

    async def test_full_paid_capacity_recalculates_and_notifies_unpaid_order(self):
        service = _orders_service()
        proof_time = datetime.datetime(2026, 8, 21)
        for user_id in range(self.shuttle_capacity):
            service.food_db.documents.append({
                "_id": ObjectId(),
                "user_id": user_id,
                "event_key": "grodno_26",
                "created_at": proof_time - datetime.timedelta(days=1),
                "proof_file": f"proof-{user_id}",
                "proof_received": proof_time,
                "choice": {
                    "total": 65,
                    "extras": {"total": 65, "shuttle": 65},
                },
            })
        unpaid_id = ObjectId()
        service.food_db.documents.append({
            "_id": unpaid_id,
            "user_id": 999,
            "event_key": "grodno_26",
            "created_at": proof_time,
            "choice": {
                "total": 165,
                "extras": {"total": 65, "shuttle": 65},
            },
        })
        reply = AsyncMock()
        service.create_update_from_user = AsyncMock(
            return_value=self._notification_update(reply)
        )

        await service.reconcile_capacity()

        unpaid = await service.food_db.find_one({"_id": unpaid_id})
        self.assertNotIn("shuttle", unpaid["choice"]["extras"])
        self.assertEqual(unpaid["choice"]["extras"]["total"], 0)
        self.assertEqual(unpaid["choice"]["total"], 0)
        self.assertTrue(unpaid["capacity_notice"]["sent"])
        reply.assert_awaited_once()

    async def test_rejected_or_cancelled_proof_releases_capacity(self):
        service = _orders_service()
        order_id = ObjectId()
        order = {
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "proof_file": "proof",
            "choice": {
                "total": 65,
                "extras": {"total": 65, "shuttle": 65},
            },
        }
        service.food_db.documents.append(order)
        await service.activate_paid_order_capacity(order_id)

        await service.release_order_capacity(order)

        self.assertTrue(await service.shuttle_available())

    async def test_proof_submitted_after_capacity_is_full_is_recalculated(self):
        service = _orders_service()
        for _ in range(self.shuttle_capacity):
            self.assertTrue(await service.reserve_shuttle_seat(ObjectId()))
        order_id = ObjectId()
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "proof_file": "proof",
            "choice": {
                "total": 165,
                "extras": {"total": 65, "shuttle": 65},
            },
        })

        order, removed = await service.activate_paid_order_capacity(order_id)

        self.assertEqual(removed, ["shuttle"])
        self.assertEqual(order["choice"]["total"], 0)
        self.assertEqual(order["capacity_notice"]["reason"], "proof_too_late")

    async def test_capacity_notice_is_claimed_before_sending(self):
        service = _orders_service()
        order = {
            "_id": ObjectId(),
            "user_id": 123,
            "event_key": "grodno_26",
            "choice": {"total": 100, "extras": {}},
            "capacity_notice": {
                "services": ["shuttle"],
                "new_total": 100,
                "reason": "unpaid",
                "sent": False,
                "sending_at": (
                    datetime.datetime.now()
                    - orders_module.NOTIFICATION_CLAIM_TTL
                    - datetime.timedelta(minutes=1)
                ),
            },
        }
        service.food_db.documents.append(order)
        reply = AsyncMock()
        service.create_update_from_user = AsyncMock(
            return_value=self._notification_update(reply)
        )

        await asyncio.gather(
            service._send_capacity_notice(order),
            service._send_capacity_notice(order),
        )

        reply.assert_awaited_once()
        current = await service.food_db.find_one({"_id": order["_id"]})
        self.assertTrue(current["capacity_notice"]["sent"])
        self.assertNotIn("sending_at", current["capacity_notice"])

    async def test_malformed_order_does_not_block_other_reconciliation(self):
        service = _orders_service()
        for _ in range(self.shuttle_capacity):
            self.assertTrue(await service.reserve_shuttle_seat(ObjectId()))
        malformed_id = ObjectId()
        valid_id = ObjectId()
        service.food_db.documents.extend([
            {
                "_id": malformed_id,
                "user_id": 1,
                "event_key": "grodno_26",
                "choice": {
                    "total": "invalid",
                    "extras": {"total": 65, "shuttle": 65},
                },
            },
            {
                "_id": valid_id,
                "user_id": 2,
                "event_key": "grodno_26",
                "choice": {
                    "total": 165,
                    "extras": {"total": 65, "shuttle": 65},
                },
            },
        ])
        reply = AsyncMock()
        service.create_update_from_user = AsyncMock(
            return_value=self._notification_update(reply)
        )

        with self.assertLogs(orders_module.logger, level="ERROR"):
            await service.reconcile_capacity()

        malformed = await service.food_db.find_one({"_id": malformed_id})
        valid = await service.food_db.find_one({"_id": valid_id})
        self.assertIn("shuttle", malformed["choice"]["extras"])
        self.assertNotIn("shuttle", valid["choice"]["extras"])

    def test_recalculation_rebuilds_forged_meal_and_extra_totals(self):
        service = _orders_service()
        choice, removed = orders_module.choice_without_capacity_services(
            {
                "total": 9999,
                "days": {
                    "friday": {
                        "total": -1,
                        "mealtimes": {
                            "dinner": {
                                "total": 5000,
                                "dishes": [{
                                    "name": "caesar",
                                    "count": 1,
                                    "price": 5000,
                                    "total": 5000,
                                }],
                                "service": {"items": [], "total": 0},
                            },
                        },
                    },
                },
                "extras": {"total": 4999, "shuttle": -1000},
            },
            {"shuttle"},
            service.menu,
        )

        self.assertEqual(removed, ["shuttle"])
        self.assertEqual(choice["days"]["friday"]["total"], 11.8)
        self.assertEqual(choice["extras"]["total"], 0)
        self.assertEqual(choice["total"], 11.8)

    def test_recalculation_supports_dishes_removed_from_current_menu(self):
        service = _orders_service()
        choice, removed = orders_module.choice_without_capacity_services(
            {
                "total": 67,
                "days": {
                    "friday": {"mealtimes": {"dinner": {"dishes": [{
                        "name": "orange_apple_juice",
                        "count": 1,
                        "price": 2,
                        "total": 2,
                    }]}}},
                },
                "extras": {"total": 65, "shuttle": 65},
            },
            {"shuttle"},
            service.menu,
        )

        self.assertEqual(removed, ["shuttle"])
        self.assertEqual(choice["extras"]["total"], 0)
        self.assertEqual(choice["total"], 2)


class PaymentAttemptTests(CapacityTestCase):
    @staticmethod
    def _admin_update(service, admin_id=999):
        service.base_app = SimpleNamespace(
            users_collection=_Collection(),
            localization=lambda key, args=None, locale=None: f"{key}:{args}",
        )
        update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        update.base = service
        update.user = admin_id
        update.bot = 10
        update.config = SimpleNamespace(
            event_key="grodno_26", admins={admin_id}, payment_admin_ru=-1
        )
        update.l = lambda key, **kwargs: f"{key}:{kwargs}"
        update.handle_cq_start = AsyncMock()
        return update

    async def test_old_accept_callback_cannot_restore_cancelled_proof(self):
        service = _orders_service()
        order_id = ObjectId()
        attempt_token = "old-attempt"
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "proof_file": "proof",
            "payment_attempt_token": attempt_token,
            "choice": {"total": 65, "extras": {"total": 65, "shuttle": 65}},
        })
        await service.activate_paid_order_capacity(order_id)

        user_update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        user_update.base = service
        user_update.user = 123
        user_update.config = SimpleNamespace(event_key="grodno_26")
        user_update.l = lambda key, **_kwargs: key
        user_update.update = SimpleNamespace(edit_or_reply=AsyncMock())
        await user_update.handle_cq_pcancel(str(order_id), attempt_token)

        admin_update = self._admin_update(service)
        await admin_update.handle_cq_adm_acc(str(order_id), attempt_token)

        order = await service.food_db.find_one({"_id": order_id})
        self.assertNotIn("proof_file", order)
        self.assertIsNot(order.get("validation"), True)
        self.assertIsNone(await service.capacity_db.find_one({
            "reservation_id": order_id
        }))
        admin_update.handle_cq_start.assert_awaited_once()

    async def test_old_reject_callback_cannot_clear_newer_cash_attempt(self):
        service = _orders_service()
        order_id = ObjectId()
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "cash_requested_at": datetime.datetime.now(),
            "payment_attempt_token": "new-attempt",
            "choice": {"total": 65, "extras": {"total": 65, "shuttle": 65}},
        })
        admin_update = self._admin_update(service)

        await admin_update.handle_cq_adm_rej(str(order_id), "old-attempt")

        order = await service.food_db.find_one({"_id": order_id})
        self.assertEqual(order["payment_attempt_token"], "new-attempt")
        self.assertIn("cash_requested_at", order)
        admin_update.handle_cq_start.assert_awaited_once()

    async def test_edit_invalidates_pending_cash_attempt(self):
        service = _orders_service()
        order_id = ObjectId()
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "cash_requested_at": datetime.datetime.now(),
            "payment_attempt_token": "cash-attempt",
            "proof_admin": 999,
            "proof_country": "be",
            "choice": {"total": 65, "extras": {"total": 65, "shuttle": 65}},
        })
        user_update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        user_update.base = service
        user_update.user = 123
        user_update.config = SimpleNamespace(event_key="grodno_26")
        user_update.handle_cq_start = AsyncMock()

        await user_update.set_choice(order_id, {"total": 5000, "extras": {}})
        admin_update = self._admin_update(service)
        await admin_update.handle_cq_adm_acc(str(order_id), "cash-attempt")

        order = await service.food_db.find_one({"_id": order_id})
        self.assertEqual(order["choice"]["total"], 0)
        self.assertNotIn("payment_attempt_token", order)
        self.assertNotIn("cash_requested_at", order)
        self.assertIsNot(order.get("validation"), True)

    async def test_legacy_cash_can_start_and_accept_new_cash_attempt(self):
        service = _orders_service()
        order_id = ObjectId()
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "proof_file": "cash",
            "choice": {
                "customer": "Alice",
                "total": 65,
                "extras": {"total": 65, "shuttle": 65},
            },
        })
        service.base_app = SimpleNamespace(
            users_collection=_Collection([
                {
                    "user_id": 999,
                    "bot_id": 10,
                    "first_name": "Admin",
                    "language_code": "en",
                    "payment_administrator_belarus": "Grodno",
                },
                {"user_id": 123, "bot_id": 10, "language_code": "en"},
            ]),
            localization=lambda key, args=None, locale=None: f"{key}:{args}",
        )
        user_update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        user_update.base = service
        user_update.user = 123
        user_update.bot = 10
        user_update.config = SimpleNamespace(event_key="grodno_26")
        user_update.l = lambda key, **kwargs: f"{key}:{kwargs}"
        user_update.update = SimpleNamespace(
            get_user=AsyncMock(return_value={"user_id": 123, "first_name": "Alice"}),
            reply=AsyncMock(),
            edit_or_reply=AsyncMock(),
        )
        user_update.handle_cq_start = AsyncMock()
        with patch.object(orders_module, "client_user_link_html", return_value="User"):
            await user_update.handle_cq_cash(str(order_id), "999")
        pending = await service.food_db.find_one({"_id": order_id})
        attempt_token = pending["payment_attempt_token"]
        self.assertNotIn("proof_file", pending)

        admin_update = self._admin_update(service)
        service.base_app.users_collection.documents.extend([
            {"user_id": 123, "bot_id": 10, "language_code": "en"},
        ])
        admin_update.update = SimpleNamespace(
            reply=AsyncMock(), edit_message_text=AsyncMock()
        )
        with patch.object(orders_module, "client_user_link_html", return_value="User"):
            await admin_update.handle_cq_adm_acc(str(order_id), attempt_token)

        accepted = await service.food_db.find_one({"_id": order_id})
        self.assertTrue(accepted["validation"])
        self.assertTrue(orders_module.order_has_payment_proof(accepted))

    async def test_stale_activation_releases_seat_after_proof_is_cancelled(self):
        service = _orders_service()
        order_id = ObjectId()
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "proof_file": "proof",
            "payment_attempt_token": "attempt",
            "choice": {"total": 65, "extras": {"total": 65, "shuttle": 65}},
        })
        reserve = service.reserve_service_seat

        async def cancel_then_reserve(
            service_name, reserved_order_id, order_snapshot=None
        ):
            await service.food_db.update_one(
                {"_id": reserved_order_id},
                {"$unset": {"proof_file": "", "payment_attempt_token": ""}},
            )
            return await reserve(service_name, reserved_order_id, order_snapshot)

        service.reserve_service_seat = cancel_then_reserve
        order, removed = await service.activate_paid_order_capacity(order_id)

        self.assertFalse(orders_module.order_has_payment_proof(order))
        self.assertEqual(removed, [])
        self.assertIsNone(await service.capacity_db.find_one({
            "reservation_id": order_id
        }))

    async def test_legacy_proof_callback_without_token_can_be_accepted(self):
        service = _orders_service()
        order_id = ObjectId()
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "proof_file": "legacy-proof",
            "choice": {"customer": "Alice", "total": 0, "extras": {}},
        })
        admin_update = self._admin_update(service)
        service.base_app.users_collection.documents.append({
            "user_id": 123, "bot_id": 10, "language_code": "en"
        })
        admin_update.update = SimpleNamespace(
            reply=AsyncMock(), edit_message_text=AsyncMock()
        )
        with patch.object(orders_module, "client_user_link_html", return_value="User"):
            await admin_update.handle_cq_adm_acc(str(order_id))

        order = await service.food_db.find_one({"_id": order_id})
        self.assertTrue(order["validation"])

    async def test_old_cancel_callback_cannot_remove_newer_proof_or_seat(self):
        service = _orders_service()
        order_id = ObjectId()
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "proof_file": "proof-b",
            "payment_attempt_token": "attempt-b",
            "choice": {"total": 65, "extras": {"total": 65, "shuttle": 65}},
        })
        await service.activate_paid_order_capacity(order_id)
        update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        update.base = service
        update.user = 123
        update.config = SimpleNamespace(event_key="grodno_26")
        update.l = lambda key, **_kwargs: key
        update.update = SimpleNamespace(edit_or_reply=AsyncMock())

        await update.handle_cq_pcancel(str(order_id), "attempt-a")

        order = await service.food_db.find_one({"_id": order_id})
        self.assertEqual(order["proof_file"], "proof-b")
        slot = await service.capacity_db.find_one({"reservation_id": order_id})
        self.assertEqual(slot["reservation_attempt_token"], "attempt-b")

    async def test_old_rejection_cannot_release_newer_attempt_reservation(self):
        service = _orders_service()
        order_id = ObjectId()
        attempt_a_time = datetime.datetime(2026, 8, 23, 10, 0)
        attempt_b_time = attempt_a_time + datetime.timedelta(minutes=1)
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "proof_file": "proof-a",
            "payment_attempt_token": "attempt-a",
            "payment_attempt_created_at": attempt_a_time,
            "choice": {
                "customer": "Alice",
                "total": 65,
                "extras": {"total": 65, "shuttle": 65},
            },
        })
        await service.activate_paid_order_capacity(order_id)
        release_attempt_a = service.release_order_capacity

        async def upload_b_before_releasing_a(order_a):
            await service.food_db.update_one({"_id": order_id}, {"$set": {
                "proof_file": "proof-b",
                "payment_attempt_token": "attempt-b",
                "payment_attempt_created_at": attempt_b_time,
            }})
            await service.activate_paid_order_capacity(order_id)
            await release_attempt_a(order_a)

        service.release_order_capacity = upload_b_before_releasing_a
        admin_update = self._admin_update(service)
        service.base_app.users_collection.documents.append({
            "user_id": 123, "bot_id": 10, "language_code": "en"
        })
        admin_update.update = SimpleNamespace(
            reply=AsyncMock(), edit_message_text=AsyncMock()
        )
        with patch.object(orders_module, "client_user_link_html", return_value="User"):
            await admin_update.handle_cq_adm_rej(str(order_id), "attempt-a")

        order = await service.food_db.find_one({"_id": order_id})
        self.assertEqual(order["payment_attempt_token"], "attempt-b")
        self.assertEqual(order["proof_file"], "proof-b")
        slot = await service.capacity_db.find_one({"reservation_id": order_id})
        self.assertEqual(slot["reservation_attempt_token"], "attempt-b")

    async def test_paid_callback_cannot_forward_another_users_proof(self):
        service = _orders_service()
        order_id = ObjectId()
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "proof_file": "proof",
            "payment_attempt_token": "attempt",
            "choice": {"customer": "Alice", "total": 0, "extras": {}},
        })
        service.base_app = SimpleNamespace(users_collection=_Collection())
        update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        update.base = service
        update.user = 456
        update.bot = 10
        update.config = SimpleNamespace(event_key="grodno_26")
        update.handle_cq_start = AsyncMock()
        update.update = SimpleNamespace(
            forward_message=AsyncMock(), reply=AsyncMock()
        )

        await update.handle_cq_paid(str(order_id))

        update.handle_cq_start.assert_awaited_once()
        update.update.forward_message.assert_not_awaited()
        order = await service.food_db.find_one({"_id": order_id})
        self.assertNotIn("proof_country", order)

    async def test_stale_paid_callback_cannot_change_cash_country(self):
        service = _orders_service()
        order_id = ObjectId()
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "proof_country": "be",
            "cash_requested_at": datetime.datetime.now(),
            "payment_attempt_token": "cash-attempt",
            "choice": {"customer": "Alice", "total": 0, "extras": {}},
        })
        update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        update.base = service
        update.user = 123
        update.config = SimpleNamespace(event_key="grodno_26")
        update.handle_cq_start = AsyncMock()

        await update.handle_cq_paid(str(order_id))

        order = await service.food_db.find_one({"_id": order_id})
        self.assertEqual(order["proof_country"], "be")

    async def test_cancelled_pdf_does_not_revive_old_legacy_cash_callback(self):
        service = _orders_service()
        order_id = ObjectId()
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "cash_requested_at": datetime.datetime(2026, 8, 23, 10, 0),
            "payment_attempt_token": "cash-attempt",
            "proof_admin": 999,
            "proof_country": "be",
            "choice": {"customer": "Alice", "total": 0, "extras": {}},
        })
        update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        update.base = service
        update.user = 123
        update.config = SimpleNamespace(event_key="grodno_26")
        update.l = lambda key, **_kwargs: key
        update.update = SimpleNamespace(
            message=SimpleNamespace(document=SimpleNamespace(file_id="pdf-proof")),
            chat_id=123,
            user=123,
            message_id=456,
            edit_or_reply=AsyncMock(),
        )

        await update.handle_payment()
        proof = await service.food_db.find_one({"_id": order_id})
        proof_token = proof["payment_attempt_token"]
        self.assertNotIn("cash_requested_at", proof)
        await update.handle_cq_pcancel(str(order_id), proof_token)

        cancelled = await service.food_db.find_one({"_id": order_id})
        self.assertFalse(_matches(
            cancelled,
            orders_module.payment_attempt_filter(
                str(order_id), None, "grodno_26"
            ),
        ))

    async def test_stale_attempt_cannot_adopt_newer_reservation(self):
        attempt_a_time = datetime.datetime(2026, 8, 23, 10, 0)
        attempt_b_time = attempt_a_time + datetime.timedelta(minutes=1)
        order_id = ObjectId()
        order_a = {
            "_id": order_id,
            "user_id": 123,
            "event_key": "grodno_26",
            "proof_file": "proof-a",
            "payment_attempt_token": "attempt-a",
            "payment_attempt_created_at": attempt_a_time,
            "choice": {"total": 65, "extras": {"total": 65, "shuttle": 65}},
        }

        class SwitchToAttemptBCollection(_Collection):
            async def find_one(self, query, projection=None):
                result = await super().find_one(query, projection)
                if query.get("payment_attempt_token") == "attempt-a":
                    self.documents[0].update({
                        "proof_file": "proof-b",
                        "payment_attempt_token": "attempt-b",
                        "payment_attempt_created_at": attempt_b_time,
                    })
                return result

        service = _orders_service()
        service.food_db = SwitchToAttemptBCollection([order_a])
        service.capacity_db.documents.append({
            "_id": "grodno_26:shuttle:0",
            "event_key": "grodno_26",
            "service": "shuttle",
            "seat": 0,
            "reservation_id": order_id,
            "reservation_attempt_token": "attempt-b",
            "reservation_attempt_created_at": attempt_b_time,
        })
        service._capacity_slots_ready.add(("grodno_26", "shuttle"))

        reserved = await service.reserve_service_seat(
            "shuttle", order_id, order_a
        )
        await service.release_order_capacity(order_a)

        self.assertFalse(reserved)
        slot = await service.capacity_db.find_one({"reservation_id": order_id})
        self.assertEqual(slot["reservation_attempt_token"], "attempt-b")

    async def test_earlier_proof_displaces_later_proof_from_last_seat(self):
        service = _orders_service()
        capacity_service = orders_module.GRODNO_OVERVIEW_SERVICE
        await service._ensure_capacity_slots(capacity_service)
        for _ in range(self.capacity_limits[capacity_service] - 1):
            self.assertTrue(await service.reserve_service_seat(
                capacity_service, ObjectId()
            ))

        earlier_time = datetime.datetime(2026, 8, 23, 10, 0)
        later_time = earlier_time + datetime.timedelta(minutes=1)
        earlier_id = ObjectId()
        later_id = ObjectId()
        for order_id, proof_time, token in (
            (earlier_id, earlier_time, "earlier"),
            (later_id, later_time, "later"),
        ):
            service.food_db.documents.append({
                "_id": order_id,
                "user_id": int(str(order_id)[-4:], 16),
                "event_key": "grodno_26",
                "created_at": earlier_time - datetime.timedelta(days=1),
                "proof_file": f"proof-{token}",
                "proof_received": proof_time,
                "payment_attempt_token": token,
                "payment_attempt_created_at": proof_time,
                "choice": {
                    "total": 25,
                    "extras": {"total": 25, capacity_service: 25},
                },
            })

        _later, later_removed = await service.activate_paid_order_capacity(later_id)
        earlier, earlier_removed = await service.activate_paid_order_capacity(earlier_id)
        later, later_removed_after = await service.activate_paid_order_capacity(later_id)

        self.assertEqual(later_removed, [])
        self.assertEqual(earlier_removed, [])
        self.assertIn(capacity_service, earlier["choice"]["extras"])
        self.assertEqual(later_removed_after, [capacity_service])
        self.assertNotIn(capacity_service, later["choice"]["extras"])
        slot = await service.capacity_db.find_one({
            "service": capacity_service,
            "reservation_id": earlier_id,
        })
        self.assertIsNotNone(slot)

    async def test_concurrent_earlier_activations_share_displaced_seat(self):
        service = _orders_service()
        capacity_service = orders_module.GRODNO_OVERVIEW_SERVICE
        await service._ensure_capacity_slots(capacity_service)
        for _ in range(self.capacity_limits[capacity_service] - 1):
            self.assertTrue(await service.reserve_service_seat(
                capacity_service, ObjectId()
            ))

        earlier_time = datetime.datetime(2026, 8, 23, 10, 0)
        order_id = ObjectId()
        later_id = ObjectId()
        service.food_db.documents.extend([
            {
                "_id": order_id,
                "user_id": 1,
                "event_key": "grodno_26",
                "created_at": earlier_time - datetime.timedelta(days=1),
                "proof_file": "proof-earlier",
                "proof_received": earlier_time,
                "payment_attempt_token": "earlier",
                "payment_attempt_created_at": earlier_time,
                "choice": {
                    "total": 25,
                    "extras": {"total": 25, capacity_service: 25},
                },
            },
            {
                "_id": later_id,
                "user_id": 2,
                "event_key": "grodno_26",
                "created_at": earlier_time - datetime.timedelta(days=1),
                "proof_file": "proof-later",
                "proof_received": earlier_time + datetime.timedelta(minutes=1),
                "payment_attempt_token": "later",
                "payment_attempt_created_at": (
                    earlier_time + datetime.timedelta(minutes=1)
                ),
                "choice": {
                    "total": 25,
                    "extras": {"total": 25, capacity_service: 25},
                },
            },
        ])
        await service.activate_paid_order_capacity(later_id)

        results = await asyncio.gather(
            service.activate_paid_order_capacity(order_id),
            service.activate_paid_order_capacity(order_id),
        )

        self.assertEqual([removed for _order, removed in results], [[], []])
        current = await service.food_db.find_one({"_id": order_id})
        self.assertIn(capacity_service, current["choice"]["extras"])
        owned_slots = [
            slot for slot in service.capacity_db.documents
            if slot.get("reservation_id") == order_id
        ]
        self.assertEqual(len(owned_slots), 1)

    async def test_priority_displacement_outlasts_capacity_many_cas_conflicts(self):
        capacity_service = orders_module.GRODNO_OVERVIEW_SERVICE
        conflict_count = self.capacity_limits[capacity_service] + 1
        earlier_time = datetime.datetime(2026, 8, 23, 10, 0)
        earlier_id = ObjectId()

        class ConflictingSwapCollection(_Collection):
            def __init__(self, documents):
                super().__init__(documents)
                self.swap_conflicts = 0

            async def update_one(self, query, update, upsert=False):
                if (
                    update.get("$set", {}).get("reservation_id") == earlier_id
                    and query.get("reservation_id") != earlier_id
                    and self.swap_conflicts < conflict_count
                ):
                    self.swap_conflicts += 1
                    return _Result(0)
                return await super().update_one(query, update, upsert)

        service = _orders_service()
        order = {
            "_id": earlier_id,
            "user_id": 1,
            "event_key": "grodno_26",
            "created_at": earlier_time - datetime.timedelta(days=1),
            "proof_file": "proof-earlier",
            "proof_received": earlier_time,
            "payment_attempt_token": "earlier",
            "payment_attempt_created_at": earlier_time,
            "choice": {
                "total": 25,
                "extras": {"total": 25, capacity_service: 25},
            },
        }
        service.food_db.documents.append(order)
        later_slots = []
        for seat in range(self.capacity_limits[capacity_service]):
            later_slots.append({
                "_id": f"grodno_26:{capacity_service}:{seat}",
                "event_key": "grodno_26",
                "service": capacity_service,
                "seat": seat,
                "reservation_id": ObjectId(),
                "reservation_attempt_token": f"later-{seat}",
                "reservation_attempt_created_at": (
                    earlier_time + datetime.timedelta(minutes=seat + 1)
                ),
            })
        service.capacity_db = ConflictingSwapCollection(later_slots)
        service._capacity_slots_ready.add(("grodno_26", capacity_service))

        reserved = await service.reserve_service_seat(
            capacity_service, earlier_id, order
        )

        self.assertTrue(reserved)
        self.assertEqual(service.capacity_db.swap_conflicts, conflict_count)
        self.assertIsNotNone(await service.capacity_db.find_one({
            "reservation_id": earlier_id
        }))

    async def test_final_free_slot_claim_handles_concurrent_release(self):
        capacity_service = orders_module.GRODNO_OVERVIEW_SERVICE
        proof_time = datetime.datetime(2026, 8, 23, 12, 0)
        order_id = ObjectId()

        class ReleaseAfterInitialClaimsCollection(_Collection):
            def __init__(self, documents):
                super().__init__(documents)
                self.free_claim_attempts = 0

            async def find_one_and_update(self, query, update, sort=None):
                if query.get("reservation_id") == {"$exists": False}:
                    self.free_claim_attempts += 1
                    result = await super().find_one_and_update(query, update, sort)
                    if self.free_claim_attempts == 3 and result is None:
                        self.documents[0].pop("reservation_id", None)
                        self.documents[0].pop("reservation_attempt_token", None)
                        self.documents[0].pop(
                            "reservation_attempt_created_at", None
                        )
                    return result
                return await super().find_one_and_update(query, update, sort)

        service = _orders_service()
        order = {
            "_id": order_id,
            "user_id": 1,
            "event_key": "grodno_26",
            "created_at": proof_time - datetime.timedelta(days=1),
            "proof_file": "proof",
            "proof_received": proof_time,
            "payment_attempt_token": "current",
            "payment_attempt_created_at": proof_time,
            "choice": {
                "total": 25,
                "extras": {"total": 25, capacity_service: 25},
            },
        }
        service.food_db.documents.append(order)
        earlier_slots = [
            {
                "_id": f"grodno_26:{capacity_service}:{seat}",
                "event_key": "grodno_26",
                "service": capacity_service,
                "seat": seat,
                "reservation_id": ObjectId(),
                "reservation_attempt_token": f"earlier-{seat}",
                "reservation_attempt_created_at": (
                    proof_time - datetime.timedelta(minutes=seat + 1)
                ),
            }
            for seat in range(self.capacity_limits[capacity_service])
        ]
        service.capacity_db = ReleaseAfterInitialClaimsCollection(earlier_slots)
        service._capacity_slots_ready.add(("grodno_26", capacity_service))

        reserved = await service.reserve_service_seat(
            capacity_service, order_id, order
        )

        self.assertTrue(reserved)
        self.assertEqual(service.capacity_db.free_claim_attempts, 4)
        self.assertIsNotNone(await service.capacity_db.find_one({
            "reservation_id": order_id
        }))

    async def test_changed_snapshot_rechecks_new_later_reservation(self):
        capacity_service = orders_module.GRODNO_OVERVIEW_SERVICE
        proof_time = datetime.datetime(2026, 8, 23, 12, 0)
        order_id = ObjectId()
        later_id = ObjectId()

        class ReplaceAfterSnapshotCollection(_Collection):
            def __init__(self, documents):
                super().__init__(documents)
                self.reserved_scans = 0

            def find(self, query, projection=None):
                cursor = super().find(query, projection)
                if query.get("reservation_id") == {"$exists": True}:
                    self.reserved_scans += 1
                    if self.reserved_scans == 1:
                        self.documents[0].update({
                            "reservation_id": later_id,
                            "reservation_attempt_token": "later",
                            "reservation_attempt_created_at": (
                                proof_time + datetime.timedelta(minutes=1)
                            ),
                        })
                return cursor

        service = _orders_service()
        order = {
            "_id": order_id,
            "user_id": 1,
            "event_key": "grodno_26",
            "created_at": proof_time - datetime.timedelta(days=1),
            "proof_file": "proof",
            "proof_received": proof_time,
            "payment_attempt_token": "current",
            "payment_attempt_created_at": proof_time,
            "choice": {
                "total": 25,
                "extras": {"total": 25, capacity_service: 25},
            },
        }
        service.food_db.documents.append(order)
        earlier_slots = [
            {
                "_id": f"grodno_26:{capacity_service}:{seat}",
                "event_key": "grodno_26",
                "service": capacity_service,
                "seat": seat,
                "reservation_id": ObjectId(),
                "reservation_attempt_token": f"earlier-{seat}",
                "reservation_attempt_created_at": (
                    proof_time - datetime.timedelta(minutes=seat + 1)
                ),
            }
            for seat in range(self.capacity_limits[capacity_service])
        ]
        service.capacity_db = ReplaceAfterSnapshotCollection(earlier_slots)
        service._capacity_slots_ready.add(("grodno_26", capacity_service))

        reserved = await service.reserve_service_seat(
            capacity_service, order_id, order
        )

        self.assertTrue(reserved)
        self.assertGreaterEqual(service.capacity_db.reserved_scans, 2)
        self.assertIsNotNone(await service.capacity_db.find_one({
            "reservation_id": order_id
        }))
        self.assertIsNone(await service.capacity_db.find_one({
            "reservation_id": later_id
        }))


class PaymentPredicateTests(unittest.TestCase):
    def test_empty_and_null_proof_values_are_queryable_as_unpaid(self):
        for proof_file in (None, "", False, "cash"):
            order = {
                "_id": ObjectId(),
                "event_key": "grodno_26",
                "proof_file": proof_file,
            }
            self.assertFalse(orders_module.order_has_payment_proof(order))
            self.assertTrue(_matches(
                order,
                orders_module.unpaid_order_filter(event_key="grodno_26"),
            ))

    def test_illegal_xml_characters_are_rejected_and_sanitized(self):
        service = _orders_service()
        with self.assertRaises(orders_module.InvalidOrderChoiceError):
            orders_module.canonicalize_choice(
                {"customer": "Bad\x00Name", "extras": {}}, service.menu
            )
        self.assertEqual(orders_module.xlsx_safe_value("Bad\x00Name"), "BadName")


class OrderPaymentReminderTests(unittest.IsolatedAsyncioTestCase):
    async def test_unpaid_order_is_reminded_once_after_two_days(self):
        service = _orders_service()
        current_time = datetime.datetime(2026, 8, 23, 12, 0)
        due_id = ObjectId()
        service.food_db.documents.extend([
            {
                "_id": due_id,
                "user_id": 100,
                "event_key": "grodno_26",
                "created_at": current_time - datetime.timedelta(days=2, minutes=1),
                "payment_reminder_sending_at": datetime.datetime(2020, 1, 1),
                "choice": {
                    "customer": "<b>Alice</b>&",
                    "total": 100,
                    "extras": {},
                },
            },
            {
                "_id": ObjectId(),
                "user_id": 200,
                "event_key": "grodno_26",
                "created_at": current_time - datetime.timedelta(days=3),
                "proof_file": "proof",
                "choice": {"customer": "Bob", "total": 100, "extras": {}},
            },
            {
                "_id": ObjectId(),
                "user_id": 300,
                "event_key": "grodno_26",
                "created_at": current_time - datetime.timedelta(days=1),
                "choice": {"customer": "Carol", "total": 100, "extras": {}},
            },
            {
                "_id": ObjectId(),
                "user_id": 400,
                "event_key": "grodno_26",
                "created_at": current_time - datetime.timedelta(days=3),
                "proof_file": "cash",
                "choice": {"customer": "Dave", "total": 100, "extras": {}},
            },
            {
                "_id": ObjectId(),
                "user_id": 500,
                "event_key": "grodno_26",
                "created_at": current_time - datetime.timedelta(days=3),
                "choice": {"customer": "Zero", "total": 0, "extras": {}},
            },
        ])
        reply = AsyncMock()
        update = SimpleNamespace(
            l=lambda key, **kwargs: f"{key}:{kwargs}",
            update=SimpleNamespace(reply=reply),
        )
        service.create_update_from_user = AsyncMock(return_value=update)

        await service.send_due_payment_reminders(current_time)
        await service.send_due_payment_reminders(current_time + datetime.timedelta(hours=1))

        due_order = await service.food_db.find_one({"_id": due_id})
        self.assertIn("payment_reminder_sent_at", due_order)
        self.assertEqual(reply.await_count, 2)
        messages = [call.args[0] for call in reply.await_args_list]
        self.assertTrue(any("&lt;b&gt;Alice&lt;/b&gt;&amp;" in text for text in messages))

    async def test_failed_reminder_is_not_sent_twice(self):
        service = _orders_service()
        current_time = datetime.datetime(2026, 8, 23, 12, 0)
        order_id = ObjectId()
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 100,
            "event_key": "grodno_26",
            "created_at": current_time - datetime.timedelta(days=3),
            "choice": {"customer": "Alice", "total": 100, "extras": {}},
        })
        reply = AsyncMock(side_effect=RuntimeError("ambiguous Telegram failure"))
        service.create_update_from_user = AsyncMock(return_value=SimpleNamespace(
            l=lambda key, **kwargs: f"{key}:{kwargs}",
            update=SimpleNamespace(reply=reply),
        ))

        with self.assertLogs(orders_module.logger, level="ERROR"):
            await service.send_due_payment_reminders(current_time)
        await service.send_due_payment_reminders(current_time + datetime.timedelta(hours=1))

        self.assertEqual(reply.await_count, 1)
        order = await service.food_db.find_one({"_id": order_id})
        self.assertIn("payment_reminder_sent_at", order)

    async def test_recalculated_zero_total_is_rechecked_at_claim(self):
        class TotalRaceCollection(_Collection):
            async def find_one_and_update(self, query, update, sort=None):
                self.documents[0]["choice"]["total"] = 0
                return await super().find_one_and_update(query, update, sort)

        service = _orders_service()
        current_time = datetime.datetime(2026, 8, 23, 12, 0)
        service.food_db = TotalRaceCollection([{
            "_id": ObjectId(),
            "user_id": 100,
            "event_key": "grodno_26",
            "created_at": current_time - datetime.timedelta(days=3),
            "choice": {"customer": "Alice", "total": 100, "extras": {}},
        }])
        service.create_update_from_user = AsyncMock()

        await service.send_due_payment_reminders(current_time)

        service.create_update_from_user.assert_not_awaited()
        order = service.food_db.documents[0]
        self.assertNotIn("payment_reminder_sent_at", order)


class CashPaymentCapacityTests(unittest.IsolatedAsyncioTestCase):
    async def test_requesting_cash_payment_does_not_count_as_paid(self):
        service = _orders_service()
        order_id = ObjectId()
        service.food_db.documents.append({
            "_id": order_id,
            "user_id": 100,
            "event_key": "grodno_26",
            "created_at": datetime.datetime(2026, 8, 20),
            "choice": {
                "customer": "<b>Alice</b>&",
                "total": 65,
                "extras": {"total": 65, "shuttle": 65},
            },
        })
        service.base_app = SimpleNamespace(
            users_collection=_Collection([{
                "user_id": 200,
                "bot_id": 10,
                "first_name": "Admin",
                "language_code": "en",
                "payment_administrator_belarus": "Grodno",
            }]),
            localization=lambda key, args=None, locale=None: f"{key}:{args}",
        )
        update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        update.base = service
        update.user = 100
        update.bot = 10
        update.config = SimpleNamespace(event_key="grodno_26")
        update.l = lambda key, **kwargs: f"{key}:{kwargs}"
        update.update = SimpleNamespace(
            get_user=AsyncMock(return_value={"user_id": 100, "first_name": "Alice"}),
            reply=AsyncMock(),
            edit_or_reply=AsyncMock(),
        )
        update.handle_cq_start = AsyncMock()

        with patch.object(orders_module, "client_user_link_html", return_value="User"):
            await update.handle_cq_cash(str(order_id), "200")

        order = await service.food_db.find_one({"_id": order_id})
        self.assertIn("cash_requested_at", order)
        self.assertNotIn("proof_file", order)
        self.assertFalse(orders_module.order_has_payment_proof(order))
        self.assertEqual(
            [slot for slot in service.capacity_db.documents if "reservation_id" in slot],
            [],
        )
        admin_message = update.update.reply.await_args.args[0]
        user_message = update.update.edit_or_reply.await_args.args[0]
        self.assertIn("&lt;b&gt;Alice&lt;/b&gt;&amp;", admin_message)
        self.assertIn("&lt;b&gt;Alice&lt;/b&gt;&amp;", user_message)


class OrdersExportTests(unittest.IsolatedAsyncioTestCase):
    async def test_mixed_sheets_are_filterable_and_summaries_are_paid_only(self):
        service = _orders_service()
        meal = {
            "dishes": [
                {"name": "dish", "count": 1, "price": 10},
                {
                    "name": "unknown",
                    "count": '=WEBSERVICE("https://example.com")',
                    "price": 0,
                },
            ],
            "service": {
                "items": [{"name": "fork", "count": 1, "price": 0.3}],
            },
        }
        for user_id, proof_file in ((1, "proof"), (2, None), (3, "cash")):
            order = {
                "_id": ObjectId(),
                "user_id": user_id,
                "event_key": "grodno_26",
                "created_at": datetime.datetime(2026, 8, 20),
                "choice": {
                    "customer": (
                        '=HYPERLINK("https://example.com")'
                        if user_id == 2
                        else ("Bad\x00Name" if user_id == 3 else f"User {user_id}")
                    ),
                    "total": 75.3,
                    "days": {
                        "friday": {"mealtimes": {"dinner": copy.deepcopy(meal)}}
                    },
                    "extras": {"total": 65, "shuttle": 65},
                },
            }
            if proof_file:
                order["proof_file"] = proof_file
            service.food_db.documents.append(order)

        captured = {}

        async def capture_workbook(_chat_id, document, **_kwargs):
            captured["workbook"] = load_workbook(document, data_only=True)

        update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        update.base = SimpleNamespace(
            food_db=service.food_db,
            menu={
                "dishes": {"dish": {"name_ru": "Блюдо"}},
                "service_items": {"fork": {"name_ru": "Вилка"}},
            },
            base_app=SimpleNamespace(users_collection=_Collection()),
        )
        update.user = 1
        update.bot = 10
        update.config = SimpleNamespace(
            admins={1}, payment_admin_ru=-1, event_key="grodno_26"
        )
        update.update = SimpleNamespace(
            user=1,
            bot=SimpleNamespace(send_document=AsyncMock(side_effect=capture_workbook)),
        )
        update.handle_cq_start = AsyncMock()

        await update.handle_cq_xlsx()

        workbook = captured["workbook"]
        self.assertEqual(
            workbook.sheetnames,
            ["Заказы", "Содержимое", "Итоги", "Посуда", "Активности"],
        )
        orders_sheet = workbook["Заказы"]
        orders_headers = [cell.value for cell in orders_sheet[1]]
        paid_column = orders_headers.index("Оплачено") + 1
        customer_column = orders_headers.index("Клиент") + 1
        self.assertEqual(
            [orders_sheet.cell(row, paid_column).value for row in (2, 3, 4)],
            [True, False, False],
        )
        formula_cell = orders_sheet.cell(3, customer_column)
        self.assertTrue(formula_cell.value.startswith("'=HYPERLINK"))
        self.assertEqual(formula_cell.data_type, "s")
        self.assertEqual(orders_sheet.cell(4, customer_column).value, "BadName")
        self.assertIsNotNone(orders_sheet.auto_filter.ref)

        details_sheet = workbook["Содержимое"]
        self.assertIsNotNone(details_sheet.auto_filter.ref)
        details_paid = {
            details_sheet.cell(row, 7).value
            for row in range(2, details_sheet.max_row + 1)
        }
        self.assertEqual(details_paid, {True, False})
        formula_counts = [
            details_sheet.cell(row, 8)
            for row in range(2, details_sheet.max_row + 1)
            if details_sheet.cell(row, 6).value == "unknown"
        ]
        self.assertTrue(formula_counts)
        self.assertTrue(all(cell.value.startswith("'=WEBSERVICE") for cell in formula_counts))
        self.assertTrue(all(cell.data_type == "s" for cell in formula_counts))

        self.assertEqual(workbook["Итоги"]["B2"].value, 1)
        self.assertEqual(workbook["Посуда"]["B2"].value, 1)
        activities = {
            row[0].value: row[1].value
            for row in workbook["Активности"].iter_rows(min_row=2)
        }
        self.assertEqual(activities["Трансфер Минск–Гродно"], 1)


class OrderEventScopingTests(unittest.IsolatedAsyncioTestCase):
    async def test_order_lookup_excludes_previous_events(self):
        service = _orders_service()
        old_order_id = ObjectId()
        current_order_id = ObjectId()
        service.food_db.documents.extend([
            {"_id": old_order_id, "user_id": 123, "event_number": 10},
            {"_id": current_order_id, "user_id": 123, "event_key": "grodno_26"},
        ])

        self.assertIsNone(await service.order_by_id(old_order_id))
        current_order = await service.order_by_id(current_order_id)
        self.assertEqual(current_order["event_key"], "grodno_26")

    async def test_new_order_is_tagged_with_current_event(self):
        service = _orders_service()
        update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        update.base = service
        update.user = 123
        update.config = SimpleNamespace(event_key="grodno_26")
        update.handle_cq_start = AsyncMock()

        await update.create_order({"extras": {}})

        self.assertEqual(len(service.food_db.documents), 1)
        self.assertEqual(service.food_db.documents[0]["event_key"], "grodno_26")
        self.assertNotIn("event_number", service.food_db.documents[0])

class GrodnoExcursionCapacityTests(CapacityTestCase):
    async def test_overview_tour_respects_capacity(self):
        service = _orders_service()
        capacity = self.capacity_limits[orders_module.GRODNO_OVERVIEW_SERVICE]
        order_ids = [ObjectId() for _ in range(capacity + 1)]

        results = await asyncio.gather(*(
            service.reserve_service_seat(
                orders_module.GRODNO_OVERVIEW_SERVICE,
                order_id,
            )
            for order_id in order_ids
        ))

        self.assertEqual(sum(results), capacity)
        self.assertFalse(await service.service_available(
            orders_module.GRODNO_OVERVIEW_SERVICE
        ))

    async def test_gorodnitsa_tour_respects_capacity(self):
        service = _orders_service()
        capacity = self.capacity_limits[orders_module.GRODNO_GORODNITSA_SERVICE]
        order_ids = [ObjectId() for _ in range(capacity + 1)]

        results = await asyncio.gather(*(
            service.reserve_service_seat(
                orders_module.GRODNO_GORODNITSA_SERVICE,
                order_id,
            )
            for order_id in order_ids
        ))

        self.assertEqual(sum(results), capacity)
        self.assertFalse(await service.service_available(
            orders_module.GRODNO_GORODNITSA_SERVICE
        ))

    async def test_full_overview_tour_rejects_order_without_saving_it(self):
        service = _orders_service()
        for _ in range(self.capacity_limits[orders_module.GRODNO_OVERVIEW_SERVICE]):
            self.assertTrue(await service.reserve_service_seat(
                orders_module.GRODNO_OVERVIEW_SERVICE,
                ObjectId(),
            ))
        update = orders_module.OrdersUpdate.__new__(orders_module.OrdersUpdate)
        update.base = service
        update.user = 123
        update.config = SimpleNamespace(event_key="grodno_26")
        choice = {"extras": {orders_module.GRODNO_OVERVIEW_SERVICE: 25}}

        with self.assertRaises(orders_module.CapacityFullError) as context:
            await update.create_order(choice)

        self.assertEqual(context.exception.service, orders_module.GRODNO_OVERVIEW_SERVICE)
        self.assertEqual(service.food_db.documents, [])

    def test_two_grodno_variants_are_rejected(self):
        choice = {"extras": {
            orders_module.GRODNO_OVERVIEW_SERVICE: 25,
            orders_module.GRODNO_GORODNITSA_SERVICE: 25,
        }}

        with self.assertRaises(orders_module.InvalidExcursionChoiceError):
            orders_module.validate_excursion_choice(choice)


if __name__ == "__main__":
    unittest.main()
